from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from config.settings import CASE_STUDY_DIR


def slugify(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", value.lower()).strip("_")


def normalize_header(value: object) -> str:
    text = str(value or "").strip().lower()
    text = text.replace("%", "pct")
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def coerce_numeric(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce")


def is_reach_sheet(headers: list[str]) -> bool:
    required = {
        "brand",
        "cumulative_reach",
        "reach_pct",
        "broadcast_spend",
        "cumlative_spend",
        "increased_reach_points",
        "cost_per_reach_point",
    }
    return required.issubset(set(headers))


def is_race_summary_sheet(ws) -> bool:
    left_header = str(ws["A1"].value or "").strip().lower()
    right_header = str(ws["J1"].value or "").strip().lower()
    return ("canidate" in left_header or "candidate" in left_header) and "primary results" in right_header


def parse_reach_sheet(ws) -> tuple[pd.DataFrame, dict[str, object]] | None:
    headers = [normalize_header(ws.cell(1, col).value) for col in range(1, ws.max_column + 1)]
    if not is_reach_sheet(headers):
        return None

    rows: list[dict[str, object]] = []
    for values in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if all(value is None for value in values[: len(headers)]):
            continue
        record = {header: values[idx] if idx < len(values) else None for idx, header in enumerate(headers)}
        brand = record.get("brand")
        day_value = record.get("date") or record.get("local_day_id")
        if brand is None and day_value is None:
            continue
        rows.append(record)

    if not rows:
        return None

    frame = pd.DataFrame(rows)
    if "local_day_id" in frame.columns and "date" not in frame.columns:
        frame.rename(columns={"local_day_id": "date"}, inplace=True)

    numeric_columns = [
        "reach",
        "cumulative_reach",
        "impressions",
        "cumulative_impressions",
        "frequency",
        "cumulative_frequency",
        "reach_pct",
        "increased_in_reach_pct",
        "broadcast_spend",
        "cumlative_spend",
        "daily_unique_hh_reached",
        "increased_reach_points",
        "cost_per_reach_point",
        "cost_per_unique_hh",
    ]
    for column in numeric_columns:
        if column in frame.columns:
            frame[column] = coerce_numeric(frame[column])

    frame["date"] = pd.to_datetime(frame["date"], errors="coerce")
    frame["sheet_name"] = ws.title
    frame["campaign_name"] = frame["brand"].fillna(ws.title)
    frame["cumulative_reach_points"] = frame["reach_pct"].fillna(0) * 100
    frame["incremental_reach_points"] = frame["increased_reach_points"].fillna(0)

    last = frame.sort_values("date").iloc[-1]
    total_broadcast_spend = float(frame["broadcast_spend"].fillna(0).sum())
    final_cumulative_spend = float(last.get("cumlative_spend") or total_broadcast_spend or 0)
    final_reach_pct = float(last.get("reach_pct") or 0)
    final_cumulative_reach = float(last.get("cumulative_reach") or 0)
    final_cost_per_reach_point = (
        final_cumulative_spend / (final_reach_pct * 100) if final_reach_pct and final_cumulative_spend else None
    )
    incremental_reach_points = float(frame["incremental_reach_points"].fillna(0).clip(lower=0).sum())
    avg_incremental_cost_per_point = (
        total_broadcast_spend / incremental_reach_points if incremental_reach_points and total_broadcast_spend else None
    )

    summary = {
        "sheet_name": ws.title,
        "campaign_name": str(last.get("campaign_name") or ws.title),
        "start_date": frame["date"].min().date().isoformat() if frame["date"].notna().any() else "",
        "end_date": frame["date"].max().date().isoformat() if frame["date"].notna().any() else "",
        "days": int(frame["date"].nunique()) if "date" in frame.columns else int(frame.shape[0]),
        "total_broadcast_spend": total_broadcast_spend,
        "final_cumulative_spend": final_cumulative_spend,
        "final_cumulative_reach": final_cumulative_reach,
        "final_reach_pct": final_reach_pct,
        "final_cumulative_frequency": float(last.get("cumulative_frequency") or 0),
        "final_cost_per_reach_point": final_cost_per_reach_point,
        "avg_incremental_cost_per_point": avg_incremental_cost_per_point,
        "final_cost_per_unique_hh": float(last.get("cost_per_unique_hh") or 0) if pd.notna(last.get("cost_per_unique_hh")) else None,
    }
    return frame, summary


def parse_race_summary_sheet(ws) -> tuple[pd.DataFrame, pd.DataFrame] | None:
    if not is_race_summary_sheet(ws):
        return None

    spend_rows: list[dict[str, object]] = []
    result_rows: list[dict[str, object]] = []
    for row in range(2, ws.max_row + 1):
        committee = ws.cell(row, 1).value
        agency = ws.cell(row, 2).value
        spend_values = [ws.cell(row, col).value for col in range(3, 9)]
        if committee is not None or agency is not None or any(value is not None for value in spend_values):
            spend_rows.append(
                {
                    "sheet_name": ws.title,
                    "candidate_committee": committee,
                    "agency_match": agency,
                    "broadcast_spend": ws.cell(row, 3).value,
                    "cable_spend": ws.cell(row, 4).value,
                    "ctv_spend": ws.cell(row, 5).value,
                    "digital_spend": ws.cell(row, 6).value,
                    "radio_spend": ws.cell(row, 7).value,
                    "grand_total": ws.cell(row, 8).value,
                }
            )

        result_candidate = ws.cell(row, 10).value
        total_ad_spend = ws.cell(row, 11).value
        vote_pct = ws.cell(row, 12).value
        if result_candidate is not None or total_ad_spend is not None or vote_pct is not None:
            result_rows.append(
                {
                    "sheet_name": ws.title,
                    "result_candidate": result_candidate,
                    "total_ad_spend": total_ad_spend,
                    "vote_pct": vote_pct,
                }
            )

    spend_frame = pd.DataFrame(spend_rows)
    results_frame = pd.DataFrame(result_rows)
    for column in [
        "broadcast_spend",
        "cable_spend",
        "ctv_spend",
        "digital_spend",
        "radio_spend",
        "grand_total",
        "total_ad_spend",
        "vote_pct",
    ]:
        if column in spend_frame.columns:
            spend_frame[column] = coerce_numeric(spend_frame[column])
        if column in results_frame.columns:
            results_frame[column] = coerce_numeric(results_frame[column])
    return spend_frame, results_frame


def build_case_study_outputs(workbook_path: Path, output_dir: Path) -> dict[str, object]:
    workbook = load_workbook(workbook_path, data_only=True)
    output_dir.mkdir(parents=True, exist_ok=True)

    reach_frames: list[pd.DataFrame] = []
    reach_summary_rows: list[dict[str, object]] = []
    spend_summary_frames: list[pd.DataFrame] = []
    result_summary_frames: list[pd.DataFrame] = []

    for ws in workbook.worksheets:
        parsed_reach = parse_reach_sheet(ws)
        if parsed_reach is not None:
            frame, summary = parsed_reach
            reach_frames.append(frame)
            reach_summary_rows.append(summary)
            continue

        parsed_summary = parse_race_summary_sheet(ws)
        if parsed_summary is not None:
            spend_frame, results_frame = parsed_summary
            if not spend_frame.empty:
                spend_summary_frames.append(spend_frame)
            if not results_frame.empty:
                result_summary_frames.append(results_frame)

    reach_daily = pd.concat(reach_frames, ignore_index=True) if reach_frames else pd.DataFrame()
    reach_summary = pd.DataFrame(reach_summary_rows)
    race_spend = pd.concat(spend_summary_frames, ignore_index=True) if spend_summary_frames else pd.DataFrame()
    race_results = pd.concat(result_summary_frames, ignore_index=True) if result_summary_frames else pd.DataFrame()

    if not reach_daily.empty:
        reach_daily.to_csv(output_dir / "reach_daily.csv", index=False)
    if not reach_summary.empty:
        reach_summary.to_csv(output_dir / "reach_summary.csv", index=False)
    if not race_spend.empty:
        race_spend.to_csv(output_dir / "race_spend_summary.csv", index=False)
    if not race_results.empty:
        race_results.to_csv(output_dir / "race_results_summary.csv", index=False)

    manifest = {
        "input_workbook": str(workbook_path),
        "output_dir": str(output_dir),
        "reach_sheets": sorted(reach_summary["sheet_name"].unique().tolist()) if not reach_summary.empty else [],
        "race_summary_sheets": sorted(race_spend["sheet_name"].dropna().unique().tolist()) if not race_spend.empty else [],
        "files_written": sorted(path.name for path in output_dir.glob("*.csv")),
    }
    (output_dir / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return manifest


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Parse a political case-study tracker workbook into reach-point and election-summary CSV outputs."
    )
    parser.add_argument("workbook", type=Path, help="Path to the source .xlsx tracker workbook.")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=CASE_STUDY_DIR / "latest",
        help="Directory where normalized CSV outputs should be written.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    manifest = build_case_study_outputs(args.workbook, args.output_dir)
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    main()
